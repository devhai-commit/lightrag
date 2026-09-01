"""
Unit tests for SpreadsheetDocumentParser.

Builds synthetic .xlsx/.csv fixtures on the fly with openpyxl/pandas into a
tmp_path, then verifies header detection, cleaning (ffill vs blank
preservation), Vietnamese number formats, truncation caps, and the
LLM-facing summary chunk's contract (must literally contain
document_id=/sheet_name= so both the model and this test can find them).
"""
from __future__ import annotations

import openpyxl
import pytest

from app.core.config import settings
from app.services.document_parser.spreadsheet_parser import SpreadsheetDocumentParser


def _parser() -> SpreadsheetDocumentParser:
    return SpreadsheetDocumentParser(workspace_id=1)


def _write_workbook(path, sheets: dict[str, list[list]]):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, data in sheets.items():
        ws = wb.create_sheet(name)
        for row in data:
            ws.append(row)
    wb.save(path)


class TestNormalCleanHeader:
    def test_clean_header_parses_correctly(self, tmp_path):
        path = tmp_path / "clean.xlsx"
        _write_workbook(path, {
            "Sheet1": [
                ["Thang", "DoanhThu"],
                ["T1", 100.5],
                ["T2", 200.25],
                ["T3", 300.75],
            ],
        })

        result = _parser().parse(str(path), document_id=1, original_filename="clean.xlsx")

        assert len(result.datasets) == 1
        dataset = result.datasets[0]
        assert dataset.sheet_name == "Sheet1"
        assert dataset.row_count == 3
        assert [c.name for c in dataset.columns] == ["Thang", "DoanhThu"]
        assert dataset.rows[0] == {"Thang": "T1", "DoanhThu": 100.5}

    def test_numeric_columns_are_native_floats_not_strings(self, tmp_path):
        path = tmp_path / "clean2.xlsx"
        _write_workbook(path, {
            "Sheet1": [["Month", "Revenue"], ["Jan", 1000], ["Feb", 2000]],
        })

        result = _parser().parse(str(path), document_id=1, original_filename="clean2.xlsx")
        dataset = result.datasets[0]
        for row in dataset.rows:
            assert isinstance(row["Revenue"], float)
            assert not isinstance(row["Revenue"], str)

    def test_tables_count_matches_sheet_count(self, tmp_path):
        path = tmp_path / "multi.xlsx"
        _write_workbook(path, {
            "A": [["X", "Y"], ["a", 1], ["b", 2]],
            "B": [["X", "Y"], ["c", 3], ["d", 4]],
        })

        result = _parser().parse(str(path), document_id=1, original_filename="multi.xlsx")
        assert result.tables_count == 2
        assert len(result.tables) == 2
        assert result.page_count == 2


class TestMessyHeaderDetection:
    def test_header_on_row_three_is_detected(self, tmp_path):
        path = tmp_path / "messy.xlsx"
        _write_workbook(path, {
            "Sheet1": [
                ["BAO CAO DOANH THU QUY 1"],
                [],
                ["Thang", "DoanhThu", "GhiChu"],
                ["T1", 100, "ok"],
                ["T2", 200, "ok"],
            ],
        })

        result = _parser().parse(str(path), document_id=1, original_filename="messy.xlsx")
        dataset = result.datasets[0]
        assert [c.name for c in dataset.columns] == ["Thang", "DoanhThu", "GhiChu"]
        assert dataset.row_count == 2
        assert dataset.rows[0]["DoanhThu"] == 100.0


class TestVietnameseNumberFormat:
    def test_vn_thousands_and_decimal_comma_parsed_correctly(self, tmp_path):
        path = tmp_path / "vn_numbers.xlsx"
        _write_workbook(path, {
            "Sheet1": [
                ["Thang", "DoanhThu"],
                ["T1", "1.234,56"],
                ["T2", "2.345,67"],
            ],
        })

        result = _parser().parse(str(path), document_id=1, original_filename="vn_numbers.xlsx")
        dataset = result.datasets[0]
        col = next(c for c in dataset.columns if c.name == "DoanhThu")
        assert col.col_type == "numeric"
        assert dataset.rows[0]["DoanhThu"] == 1234.56
        assert dataset.rows[1]["DoanhThu"] == 2345.67


class TestMergedCellFfill:
    def test_label_column_ffilled_but_numeric_column_stays_blank(self, tmp_path):
        path = tmp_path / "merged.xlsx"
        _write_workbook(path, {
            "Sheet1": [
                ["Quy", "Thang", "DoanhThu"],
                ["Q1", "T1", 100],
                [None, "T2", None],
                [None, "T3", 300],
                ["Q2", "T4", 400],
            ],
        })

        result = _parser().parse(str(path), document_id=1, original_filename="merged.xlsx")
        dataset = result.datasets[0]

        # Label column (Quy) forward-filled across the merged-cell group.
        assert [row["Quy"] for row in dataset.rows] == ["Q1", "Q1", "Q1", "Q2"]

        # Numeric column keeps its genuine blank as None — never invented.
        assert dataset.rows[1]["DoanhThu"] is None
        assert dataset.rows[0]["DoanhThu"] == 100.0
        assert dataset.rows[2]["DoanhThu"] == 300.0


class TestEmptyTrailingRowsColumns:
    def test_trailing_empty_rows_and_columns_dropped(self, tmp_path):
        path = tmp_path / "trailing.xlsx"
        _write_workbook(path, {
            "Sheet1": [
                ["Thang", "DoanhThu", None],
                ["T1", 100, None],
                ["T2", 200, None],
                [None, None, None],
                [None, None, None],
            ],
        })

        result = _parser().parse(str(path), document_id=1, original_filename="trailing.xlsx")
        dataset = result.datasets[0]
        assert dataset.row_count == 2
        assert [c.name for c in dataset.columns] == ["Thang", "DoanhThu"]


class TestRowCapTruncation:
    def test_truncation_when_exceeding_monkeypatched_cap(self, tmp_path, monkeypatch):
        monkeypatch.setattr(settings, "NEXUSRAG_SPREADSHEET_MAX_ROWS_PER_SHEET", 3)

        rows = [["Thang", "DoanhThu"]] + [[f"T{i}", i * 10] for i in range(10)]
        path = tmp_path / "big.xlsx"
        _write_workbook(path, {"Sheet1": rows})

        result = _parser().parse(str(path), document_id=1, original_filename="big.xlsx")
        dataset = result.datasets[0]

        assert dataset.row_count == 10  # true count before truncation
        assert dataset.truncated is True
        assert dataset.truncated_at_row == 3
        assert len(dataset.rows) == 3


class TestSummaryChunkContract:
    def test_summary_chunk_contains_document_id_and_sheet_name(self, tmp_path):
        path = tmp_path / "clean.xlsx"
        _write_workbook(path, {
            "Revenue Q1": [["Thang", "DoanhThu"], ["T1", 100], ["T2", 200]],
        })

        result = _parser().parse(str(path), document_id=77, original_filename="clean.xlsx")
        chunk = result.chunks[0]
        assert "document_id=77" in chunk.content
        assert 'sheet_name="Revenue Q1"' in chunk.content
        assert chunk.has_table is True
        assert chunk.page_no == 1
        assert chunk.document_id == 77


class TestCsvSupport:
    def test_csv_parsed_as_single_implicit_sheet(self, tmp_path):
        path = tmp_path / "data.csv"
        path.write_text("Thang,DoanhThu\nT1,100\nT2,200\n", encoding="utf-8")

        result = _parser().parse(str(path), document_id=1, original_filename="data.csv")
        assert len(result.datasets) == 1
        assert result.datasets[0].row_count == 2
        assert result.datasets[0].rows[0]["DoanhThu"] == 100.0


def test_supported_extensions():
    assert SpreadsheetDocumentParser.supported_extensions() == {".xlsx", ".csv"}


def test_parser_name():
    assert SpreadsheetDocumentParser.parser_name == "spreadsheet"
